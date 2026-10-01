"""
Índice de dimensões (protótipo): uma linha por cada tamanho em pixels
("NNNxNNN") encontrado no texto da coluna Dimensão dos formatos Digital
(Internet + Programático), com o contexto de onde veio (ex.: "Mobile",
"Cross-device recomendado") e o Fornecedor/Veículo/Formato de origem —
para se conseguir pesquisar/filtrar por uma dimensão concreta (ex.:
"300x600") em toda a base, em vez de ter de ler célula a célula.

Extração mecânica (regex), nunca reinterpretação do conteúdo: cada
linha do índice é um excerto literal do texto já existente na base.

Correr a partir de qualquer diretoria:
    python3 scripts/gerar_indice_dimensoes.py
"""
import os
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

RAIZ_PROJETO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ORIGEM = os.path.join(RAIZ_PROJETO, "data", "base-formatos.xlsx")
DESTINO = os.path.join(RAIZ_PROJETO, "data", "CSBuilder_Indice_Dimensoes.xlsx")

PADRAO_DIMENSAO = re.compile(r"\b\d{2,5}\s*[x×]\s*\d{2,5}\b")


def dividir_em_clausulas(texto):
    """Divide o texto em cláusulas por '.' e ';' de topo (fora de
    parênteses), para cada cláusula poder ser lida como "um facto"."""
    partes = []
    atual = []
    profundidade = 0
    for caractere in texto:
        if caractere == "(":
            profundidade += 1
        elif caractere == ")":
            profundidade -= 1
        if caractere in ".;" and profundidade == 0:
            partes.append("".join(atual))
            atual = []
        else:
            atual.append(caractere)
    partes.append("".join(atual))
    return [p.strip() for p in partes if p.strip()]


def extrair_contexto_e_dimensoes(texto):
    """Devolve uma lista [(contexto, dimensao), ...] — uma entrada por
    cada tamanho em pixels encontrado no texto, com o rótulo da
    cláusula onde apareceu (ex.: "Mobile: 1200x1200, 1080x1440" dá duas
    entradas, ambas com contexto "Mobile"). Sem rótulo reconhecível,
    usa-se a própria cláusula (truncada) como contexto."""
    if not texto:
        return []
    resultados = []
    for clausula in dividir_em_clausulas(texto):
        ocorrencias = list(PADRAO_DIMENSAO.finditer(clausula))
        if not ocorrencias:
            continue
        correspondencia_rotulo = re.match(r"^([^:]{2,50}):\s*(.*)$", clausula)
        if correspondencia_rotulo:
            contexto = correspondencia_rotulo.group(1).strip()
        else:
            contexto = clausula[:60].strip() + ("…" if len(clausula) > 60 else "")
        for ocorrencia in ocorrencias:
            dimensao = ocorrencia.group(0).replace(" ", "").replace("×", "x")
            resultados.append((contexto, dimensao))
    return resultados


wb_origem = openpyxl.load_workbook(ORIGEM, data_only=True)
ws_origem = wb_origem.active
cabecalhos = [c.value for c in ws_origem[1]]
idx = {h: i for i, h in enumerate(cabecalhos)}

linhas_indice = []
for row in ws_origem.iter_rows(min_row=2, values_only=True):
    meio = (row[idx["Meio"]] or "").upper()
    if meio not in ("INTERNET", "PROGRAMÁTICO"):
        continue
    fornecedor = row[idx["Fornecedor"]]
    veiculo = row[idx["Veículo"]]
    formato = row[idx["Formato"]]
    for contexto, dimensao in extrair_contexto_e_dimensoes(row[idx["Dimensão"]]):
        linhas_indice.append((dimensao, fornecedor, veiculo, formato, contexto))

# Ordenado pela própria dimensão primeiro — abres o ficheiro e já vês
# tudo agrupado por tamanho; o AutoFilter permite reordenar por
# Fornecedor/Formato se for mais útil num dado momento.
linhas_indice.sort(key=lambda linha: (linha[0], linha[1] or "", linha[3] or ""))

wb = openpyxl.Workbook()
folha = wb.active
folha.title = "Índice de Dimensões"
folha.freeze_panes = "A2"
folha.views.sheetView[0].showGridLines = False

COLUNAS = ["Dimensão", "Fornecedor", "Veículo", "Formato", "Contexto"]
LARGURAS = {"Dimensão": 16, "Fornecedor": 22, "Veículo": 22, "Formato": 28, "Contexto": 55}

FONTE_CABECALHO = Font(bold=True, color="FFFFFF")
FUNDO_CABECALHO = PatternFill("solid", fgColor="1A1A1A")
FONTE_CELULA = Font(size=10)
ALINHAMENTO = Alignment(vertical="center", wrap_text=True)
BORDA_CLARA = Side(style="thin", color="DCDCDC")
BORDA_COMPLETA = Border(top=BORDA_CLARA, left=BORDA_CLARA, bottom=BORDA_CLARA, right=BORDA_CLARA)

for i, nome_coluna in enumerate(COLUNAS, start=1):
    celula = folha.cell(row=1, column=i, value=nome_coluna)
    celula.font = FONTE_CABECALHO
    celula.fill = FUNDO_CABECALHO
    celula.alignment = Alignment(vertical="center")
    celula.border = BORDA_COMPLETA
    folha.column_dimensions[get_column_letter(i)].width = LARGURAS[nome_coluna]

for linha_num, linha in enumerate(linhas_indice, start=2):
    for i, valor in enumerate(linha, start=1):
        celula = folha.cell(row=linha_num, column=i, value=valor)
        celula.font = FONTE_CELULA
        celula.alignment = ALINHAMENTO
        celula.border = BORDA_COMPLETA

folha.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{len(linhas_indice) + 1}"

wb.save(DESTINO)
print("Gerado:", DESTINO)
print("Total de linhas no índice:", len(linhas_indice))
