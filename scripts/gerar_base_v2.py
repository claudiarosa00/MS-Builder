"""
CSBuilder_Base_Formatos v2: versão mais completa e normalizada da cópia
de referência da base, combinando o que foi validado nos protótipos
anteriores:

1. Placements reais (Facebook/Instagram: Single Image, Single Video,
   Carousel, Collection) saem como uma linha por placement
   (Feed/Stories/Reels/Explorar), com Dimensão/Aspect Ratio/Copies
   próprios — ver gerar_base_detalhada_placements.py.
2. Peso e Tipo de Ficheiro ganham cada um duas colunas fixas —
   "... Imagem" e "... Vídeo" — nos formatos onde o texto original já
   distinguia os dois (ex.: "Imagem: 30 MB; Vídeo: 4 GB"). Nos
   formatos com um valor só (sem essa distinção), todo o valor vai
   para a coluna "... Imagem", por omissão.
3. Cobre todos os meios (Digital, OOH, TV, Rádio, Cinema, Imprensa).

Tudo por extração mecânica do texto já existente na base — nunca
reinterpretação nem invenção de conteúdo.

Correr a partir de qualquer diretoria:
    python3 scripts/gerar_base_v2.py
"""
import os
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

RAIZ_PROJETO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ORIGEM = os.path.join(RAIZ_PROJETO, "data", "base-formatos.xlsx")
DESTINO = os.path.join(RAIZ_PROJETO, "data", "CSBuilder_Base_Formatos_v2.xlsx")

MAPA_MEIO = {
    "INTERNET": "Digital", "PROGRAMÁTICO": "Digital", "OOH": "OOH", "TV": "TV",
    "RÁDIO": "Rádio", "RADIO": "Rádio", "CINEMA": "Cinema", "IMPRENSA": "Imprensa",
}
ORDEM_MEIOS = ["Digital", "OOH", "TV", "Rádio", "Cinema", "Imprensa"]

# Mesma lista revista um a um com a Cláudia — só estes 8 formatos têm
# placements reais (sítios diferentes da mesma rede, cada um com a sua
# dimensão); todos os outros formatos com texto "Rótulo: valor" parecido
# ficam de fora de propósito (ver análise feita em conversa).
FORMATOS_COM_PLACEMENT = {("Facebook", "Single Image"), ("Facebook", "Single Video"),
                           ("Facebook", "Carousel"), ("Facebook", "Collection"),
                           ("Instagram", "Single Image"), ("Instagram", "Single Video"),
                           ("Instagram", "Carousel"), ("Instagram", "Collection")}
PLACEMENTS_CONHECIDOS = ["Feed", "Stories", "Reels", "Explorar"]


def dividir_em_clausulas(texto):
    """Divide o texto em cláusulas por '.' e ';' de topo (fora de
    parênteses). Um '.' entre dois dígitos (ex.: "1.5 MB", "40.000c" —
    a base mistura "." como separador decimal e de milhares) nunca é
    tratado como fim de frase, para não partir o número ao meio."""
    partes = []
    atual = []
    profundidade = 0
    for i, caractere in enumerate(texto):
        if caractere == "(":
            profundidade += 1
        elif caractere == ")":
            profundidade -= 1
        ponto_dentro_de_numero = (
            caractere == "."
            and i > 0 and i + 1 < len(texto)
            and texto[i - 1].isdigit() and texto[i + 1].isdigit()
        )
        if caractere in ".;" and profundidade == 0 and not ponto_dentro_de_numero:
            partes.append("".join(atual))
            atual = []
        else:
            atual.append(caractere)
    partes.append("".join(atual))
    return [p.strip() for p in partes if p.strip()]


def dividir_so_por_ponto_e_virgula(texto):
    """Como dividir_em_clausulas, mas só separa em ';' — nunca em '.'.
    Os 8 formatos do Meta usam sempre ';' entre placements, com pontos
    perfeitamente normais lá dentro (ex.: "Feed: 1080x1080 mín. (capa:
    imagem ou vídeo) + 3 imagens de produto"), que um split por '.'
    partiria ao meio."""
    partes = []
    atual = []
    profundidade = 0
    for caractere in texto:
        if caractere == "(":
            profundidade += 1
        elif caractere == ")":
            profundidade -= 1
        if caractere == ";" and profundidade == 0:
            partes.append("".join(atual))
            atual = []
        else:
            atual.append(caractere)
    partes.append("".join(atual))
    return [p.strip().rstrip(".").strip() for p in partes if p.strip()]


def repartir_por_placement(texto):
    """'Feed: a; Stories: b; Reels: c.' -> {'Feed': 'a', 'Stories': 'b', 'Reels': 'c'}."""
    if not texto:
        return {}
    resultado = {}
    for clausula in dividir_so_por_ponto_e_virgula(texto):
        correspondencia = re.match(r"^(Feed|Stories|Reels|Explorar)\s*:\s*(.*)$", clausula)
        if correspondencia:
            resultado[correspondencia.group(1)] = correspondencia.group(2).strip()
    return resultado


def dividir_imagem_video(texto):
    """Devolve (valor_imagem, valor_video). Quando o texto tem
    cláusulas "Imagem: ..." e/ou "Vídeo: ...", usa esses valores
    (outras cláusulas, ex.: "Thumbnail: ...", ficam anexadas a Imagem,
    para não se perderem). Sem nenhuma dessas duas cláusulas, o texto
    todo vai para Imagem, por omissão — Vídeo fica em branco."""
    if not texto:
        return "", ""
    partes_imagem = []
    partes_video = []
    partes_outras = []
    tem_split = False
    for clausula in dividir_em_clausulas(texto):
        correspondencia = re.match(r"^(Imagem|V[ií]deo)\s*:\s*(.*)$", clausula, re.IGNORECASE)
        if correspondencia:
            tem_split = True
            rotulo = correspondencia.group(1).lower()
            valor = correspondencia.group(2).strip()
            (partes_imagem if rotulo.startswith("imagem") else partes_video).append(valor)
        else:
            partes_outras.append(clausula)
    if not tem_split:
        return texto.strip(), ""
    imagem = "; ".join(partes_imagem + partes_outras)
    video = "; ".join(partes_video)
    return imagem, video


wb_origem = openpyxl.load_workbook(ORIGEM, data_only=True)
ws_origem = wb_origem.active
cabecalhos = [c.value for c in ws_origem[1]]
idx = {h: i for i, h in enumerate(cabecalhos)}

linhas_por_meio = {m: [] for m in ORDEM_MEIOS}
for row in ws_origem.iter_rows(min_row=2, values_only=True):
    meio_bruto = (row[idx["Meio"]] or "").upper()
    grupo = MAPA_MEIO.get(meio_bruto)
    if grupo:
        linhas_por_meio[grupo].append(row)

COLUNAS_BASE = ["Fornecedor", "Veículo", "Formato", "Placement", "Dimensão", "Aspect Ratio",
                "Peso Imagem", "Peso Vídeo", "Tipo de Ficheiro Imagem", "Tipo de Ficheiro Vídeo",
                "Observações", "Link"]
LARGURAS = {
    "Fornecedor": 22, "Veículo": 22, "Grupo Digital2020": 18, "Formato": 26, "Placement": 12,
    "Dimensão": 40, "Aspect Ratio": 18, "Peso Imagem": 16, "Peso Vídeo": 16,
    "Tipo de Ficheiro Imagem": 22, "Tipo de Ficheiro Vídeo": 22,
    "Copies": 38, "Entrega": 22, "Observações": 42, "Link": 38,
}

wb = openpyxl.Workbook()
wb.remove(wb.active)

FONTE_CABECALHO = Font(bold=True, color="FFFFFF")
FUNDO_CABECALHO = PatternFill("solid", fgColor="1A1A1A")
FONTE_CELULA = Font(size=10)
ALINHAMENTO = Alignment(vertical="center", wrap_text=True)
BORDA_CLARA = Side(style="thin", color="DCDCDC")
BORDA_COMPLETA = Border(top=BORDA_CLARA, left=BORDA_CLARA, bottom=BORDA_CLARA, right=BORDA_CLARA)
FUNDO_PLACEMENT = PatternFill("solid", fgColor="F4F7FC")

contagem_placement_gerados = 0

for meio in ORDEM_MEIOS:
    linhas = linhas_por_meio[meio]
    if not linhas:
        continue

    colunas = list(COLUNAS_BASE)
    if meio == "Digital":
        colunas.insert(2, "Grupo Digital2020")
        colunas.insert(colunas.index("Observações"), "Copies")
    if meio == "TV":
        colunas.insert(colunas.index("Observações"), "Entrega")

    folha = wb.create_sheet(meio)
    folha.freeze_panes = "A2"
    folha.views.sheetView[0].showGridLines = False

    for i, nome_coluna in enumerate(colunas, start=1):
        celula = folha.cell(row=1, column=i, value=nome_coluna)
        celula.font = FONTE_CABECALHO
        celula.fill = FUNDO_CABECALHO
        celula.alignment = Alignment(vertical="center")
        celula.border = BORDA_COMPLETA
        folha.column_dimensions[get_column_letter(i)].width = LARGURAS.get(nome_coluna, 20)

    linha_atual = 2
    for row in linhas:
        fornecedor = row[idx["Fornecedor"]]
        formato_nome = row[idx["Formato"]]
        chave = (fornecedor, formato_nome)

        peso_imagem, peso_video = dividir_imagem_video(row[idx["Peso"]])
        tipo_imagem, tipo_video = dividir_imagem_video(row[idx["Tipo de Ficheiro"]])

        if chave in FORMATOS_COM_PLACEMENT:
            dim_por_placement = repartir_por_placement(row[idx["Dimensão"]])
            ar_por_placement = repartir_por_placement(row[idx["Aspect Ratio"]])
            copies_por_placement = repartir_por_placement(row[idx["Copies"]]) if "Copies" in idx else {}
            placements_desta_linha = [p for p in PLACEMENTS_CONHECIDOS if p in dim_por_placement]
        else:
            placements_desta_linha = [None]

        for placement in placements_desta_linha:
            for i, nome_coluna in enumerate(colunas, start=1):
                celula_fill = FUNDO_PLACEMENT if placement else None
                if nome_coluna == "Placement":
                    valor = placement or ""
                elif nome_coluna == "Entrega":
                    celula = folha.cell(row=linha_atual, column=i, value="GoFastWay\nwww.gofastway.tv")
                    celula.hyperlink = "http://www.gofastway.tv"
                    celula.font = Font(size=10, color="1155CC", underline="single")
                    celula.alignment = ALINHAMENTO
                    celula.border = BORDA_COMPLETA
                    continue
                elif nome_coluna == "Link":
                    link = row[idx["Link"]]
                    if link:
                        celula = folha.cell(row=linha_atual, column=i, value=link)
                        celula.hyperlink = link
                        celula.font = Font(size=10, color="1155CC", underline="single")
                        celula.alignment = ALINHAMENTO
                        celula.border = BORDA_COMPLETA
                        if celula_fill:
                            celula.fill = celula_fill
                        continue
                    valor = None
                elif nome_coluna == "Peso Imagem":
                    valor = peso_imagem
                elif nome_coluna == "Peso Vídeo":
                    valor = peso_video
                elif nome_coluna == "Tipo de Ficheiro Imagem":
                    valor = tipo_imagem
                elif nome_coluna == "Tipo de Ficheiro Vídeo":
                    valor = tipo_video
                elif placement and nome_coluna == "Dimensão":
                    valor = dim_por_placement.get(placement, "")
                elif placement and nome_coluna == "Aspect Ratio":
                    valor = ar_por_placement.get(placement, "")
                elif placement and nome_coluna == "Copies":
                    valor = copies_por_placement.get(placement, "")
                else:
                    valor = row[idx[nome_coluna]] if nome_coluna in idx else None
                celula = folha.cell(row=linha_atual, column=i, value=valor)
                celula.font = FONTE_CELULA
                celula.alignment = ALINHAMENTO
                celula.border = BORDA_COMPLETA
                if celula_fill:
                    celula.fill = celula_fill
            linha_atual += 1
            if placement:
                contagem_placement_gerados += 1

    folha.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{linha_atual - 1}"

wb.save(DESTINO)
print("Gerado:", DESTINO)
for meio in ORDEM_MEIOS:
    if linhas_por_meio[meio]:
        print(f"  {meio}: {len(linhas_por_meio[meio])} formatos na base")
print(f"Linhas de placement geradas (Meta): {contagem_placement_gerados} (a partir de 8 formatos)")
