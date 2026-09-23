"""
Gera uma base de formatos em Excel, 1 meio por separador (Digital, OOH,
TV, Rádio, Imprensa), com exatamente as colunas que aparecem na
Biblioteca de Formatos do CSBuilder (ver criarCabecalhoTabela em app.js):
Fornecedor, Veículo, Grupo Digital2020 (só Digital), Formato, Dimensão,
Aspect Ratio, Peso, Tipo de Ficheiro, Copies (só Digital), Entrega (só
TV, sempre "GoFastWay"), Observações, Link.

Cinema fica de fora porque não existe atualmente nenhum formato desse
meio na base (nada aparece na Biblioteca de Formatos para Cinema hoje).
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Caminhos relativos a este script (scripts/), para correr a partir de
# qualquer diretoria — a base fica em data/, tal como o resto da app.
RAIZ_PROJETO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ORIGEM = os.path.join(RAIZ_PROJETO, "data", "base-formatos.xlsx")
DESTINO = os.path.join(RAIZ_PROJETO, "data", "CSBuilder_Base_Formatos.xlsx")

# Mesmo mapeamento que VALOR_BASE_PARA_GRUPO_MEIO em app.js
MAPA_MEIO = {
    "INTERNET": "Digital",
    "PROGRAMÁTICO": "Digital",
    "OOH": "OOH",
    "TV": "TV",
    "RÁDIO": "Rádio",
    "RADIO": "Rádio",
    "IMPRENSA": "Imprensa",
}
ORDEM_MEIOS = ["Digital", "OOH", "TV", "Rádio", "Imprensa"]

wb_origem = openpyxl.load_workbook(ORIGEM, data_only=True)
ws_origem = wb_origem.active
cabecalhos = [c.value for c in ws_origem[1]]
idx = {h: i for i, h in enumerate(cabecalhos)}

linhas_por_meio = {m: [] for m in ORDEM_MEIOS}
for row in ws_origem.iter_rows(min_row=2, values_only=True):
    meio_bruto = (row[idx["Meio"]] or "").upper()
    grupo = MAPA_MEIO.get(meio_bruto)
    if grupo:
        linhas_por_meio[grupo].append(row)

COLUNAS_BASE = ["Fornecedor", "Veículo", "Formato", "Dimensão", "Aspect Ratio", "Peso", "Tipo de Ficheiro", "Observações", "Link"]
LARGURAS = {
    "Fornecedor": 22, "Veículo": 22, "Grupo Digital2020": 18, "Formato": 30,
    "Dimensão": 45, "Aspect Ratio": 14, "Peso": 18, "Tipo de Ficheiro": 24,
    "Copies": 40, "Entrega": 16, "Observações": 45, "Link": 40,
}

wb = openpyxl.Workbook()
wb.remove(wb.active)

FONTE_CABECALHO = Font(bold=True, color="FFFFFF")
FUNDO_CABECALHO = PatternFill("solid", fgColor="1A1A1A")
FONTE_CELULA = Font(size=10)
ALINHAMENTO = Alignment(vertical="top", wrap_text=True)

# Grelha cinzenta clara — mesma cor/estilo já usado no Excel do pedido
# (ver ESTILO_BORDA_COMPLETA em app.js), para as duas exportações ficarem
# visualmente consistentes.
BORDA_CLARA = Side(style="thin", color="DCDCDC")
BORDA_COMPLETA = Border(top=BORDA_CLARA, left=BORDA_CLARA, bottom=BORDA_CLARA, right=BORDA_CLARA)

for meio in ORDEM_MEIOS:
    linhas = linhas_por_meio[meio]
    if not linhas:
        continue

    colunas = list(COLUNAS_BASE)
    # Formato: Fornecedor, Veículo, [Grupo Digital2020], Formato, Dimensão, ...
    if meio == "Digital":
        colunas.insert(2, "Grupo Digital2020")
        colunas.insert(colunas.index("Observações"), "Copies")
    if meio == "TV":
        colunas.insert(colunas.index("Observações"), "Entrega")

    folha = wb.create_sheet(meio)
    folha.freeze_panes = "A2"
    folha.views.sheetView[0].showGridLines = False

    for i, nome_coluna in enumerate(colunas, start=1):
        celula = folha.cell(row=1, column=i, value=nome_coluna)
        celula.font = FONTE_CABECALHO
        celula.fill = FUNDO_CABECALHO
        celula.alignment = Alignment(vertical="center")
        celula.border = BORDA_COMPLETA
        folha.column_dimensions[get_column_letter(i)].width = LARGURAS.get(nome_coluna, 20)

    linha_atual = 2
    for row in linhas:
        for i, nome_coluna in enumerate(colunas, start=1):
            if nome_coluna == "Entrega":
                valor = "GoFastWay"
            elif nome_coluna == "Link":
                link = row[idx["Link"]]
                if link:
                    celula = folha.cell(row=linha_atual, column=i, value=link)
                    celula.hyperlink = link
                    celula.font = Font(size=10, color="1155CC", underline="single")
                    celula.alignment = ALINHAMENTO
                    celula.border = BORDA_COMPLETA
                    continue
                valor = None
            else:
                valor = row[idx[nome_coluna]] if nome_coluna in idx else None
            celula = folha.cell(row=linha_atual, column=i, value=valor)
            celula.font = FONTE_CELULA
            celula.alignment = ALINHAMENTO
            celula.border = BORDA_COMPLETA
        linha_atual += 1

    folha.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{linha_atual - 1}"

wb.save(DESTINO)
print("Gerado:", DESTINO)
for meio in ORDEM_MEIOS:
    print(f"  {meio}: {len(linhas_por_meio[meio])} linhas")
